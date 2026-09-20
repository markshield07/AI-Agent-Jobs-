"""Pull plain text out of an uploaded resume.

Deliberately dumb: get the words out, let the parser find the structure. A PDF
that yields nothing is almost always a scan, so say that rather than handing the
parser an empty string to hallucinate against.
"""

from __future__ import annotations

import hashlib
import io

SUPPORTED_SUFFIXES = {".pdf", ".docx", ".txt", ".md"}


class UnsupportedResume(ValueError):
    """The file is not a format we can read text out of."""


class EmptyResume(ValueError):
    """The file parsed, but produced no usable text."""


def content_sha(data: bytes) -> str:
    return hashlib.sha256(data).hexdigest()


def suffix_of(filename: str) -> str:
    _, _, ext = filename.rpartition(".")
    return f".{ext.lower()}" if ext else ""


def extract_text(data: bytes, filename: str) -> str:
    """Return the resume's text, or raise if it cannot be read."""
    suffix = suffix_of(filename)
    if suffix not in SUPPORTED_SUFFIXES:
        raise UnsupportedResume(
            f"{filename}: expected one of {', '.join(sorted(SUPPORTED_SUFFIXES))}"
        )

    if suffix == ".pdf":
        text = _from_pdf(data)
    elif suffix == ".docx":
        text = _from_docx(data)
    else:
        text = data.decode("utf-8", errors="replace")

    text = _tidy(text)
    if not text:
        raise EmptyResume(
            f"{filename}: no text found. If this is a scanned PDF, "
            "export a text version or paste the resume as .txt."
        )
    return text


def _from_pdf(data: bytes) -> str:
    from pypdf import PdfReader

    reader = PdfReader(io.BytesIO(data))
    return "\n".join(page.extract_text() or "" for page in reader.pages)


def _from_docx(data: bytes) -> str:
    import docx

    document = docx.Document(io.BytesIO(data))
    parts = [p.text for p in document.paragraphs]
    for table in document.tables:
        for row in table.rows:
            parts.extend(cell.text for cell in row.cells)
    return "\n".join(parts)


def _tidy(text: str) -> str:
    lines = [line.rstrip() for line in text.replace("\r\n", "\n").split("\n")]
    out: list[str] = []
    for line in lines:
        if line or (out and out[-1]):
            out.append(line)
    return "\n".join(out).strip()
